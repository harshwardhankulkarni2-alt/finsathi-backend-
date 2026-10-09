"""Parses bank statement CSV / XLSX files into normalized transaction rows.

Real bank exports vary a lot, so this is deliberately tolerant: it finds the
header row (statements often start with account details), detects columns by
keyword, and supports separate Debit/Credit columns or a single Amount column.
PDF parsing is intentionally not implemented yet (see README).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from .categorize import guess_category, guess_payment_method


class ParseError(ValueError):
    pass


@dataclass
class ParsedRow:
    date: date
    merchant: str
    amount: float
    category: str
    payment_method: str
    notes: str
    needs_review: bool


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    skipped: int = 0
    warnings: list[str] = field(default_factory=list)


DATE_FORMATS = [
    "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y",
    "%d %b %Y", "%d-%b-%Y", "%d %B %Y", "%d-%b-%y", "%b %d, %Y", "%Y/%m/%d",
    "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M:%S",
]


def parse_date(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:  # Excel serial date
            return (datetime(1899, 12, 30) + timedelta(days=int(v))).date()
        return None
    s = str(v).strip()
    if not s:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_amount(v: Any) -> tuple[float, str | None]:
    """Returns (value, hint) where hint is 'dr' / 'cr' if the cell carried a Dr/Cr suffix."""
    if v is None or v == "":
        return 0.0, None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v), None
    s = str(v).strip()
    hint = None
    m = re.search(r"\b(dr|cr)\.?$", s, re.I)
    if m:
        hint = m.group(1).lower()
        s = s[: m.start()].strip()
    negative = s.startswith("(") and s.endswith(")")
    s = re.sub(r"[^\d.\-]", "", s.replace("₹", "").replace("Rs.", "").replace("Rs", ""))
    if s in ("", "-", ".", "-."):
        return 0.0, hint
    try:
        val = float(s)
    except ValueError:
        return 0.0, hint
    return (-abs(val) if negative else val), hint


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return content.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ParseError("Could not read the file encoding")


def _read_csv(content: bytes) -> list[list[Any]]:
    text = _decode(content)
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(content: bytes) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    except Exception as e:  # corrupt / not really an xlsx
        raise ParseError("Could not read this Excel file. Is it a valid .xlsx?") from e


def _norm(cell: Any) -> str:
    return re.sub(r"\s+", " ", str(cell or "")).strip().lower()


def _find(headers: list[str], keywords: list[str], exclude: tuple[str, ...] = ()) -> int:
    for kw in keywords:  # keyword priority order
        for i, h in enumerate(headers):
            if kw in h and not any(x in h for x in exclude):
                return i
    return -1


MERCHANT_KW = ["narration", "description", "particulars", "merchant", "details", "payee", "remarks"]
NOISE_TOKENS = {"upi", "cr", "dr", "p2a", "p2m", "p2p", "imps", "neft", "rtgs", "pos", "ach", "payment", "from", "to", "paid", "via"}


def clean_merchant(narration: str) -> str:
    text = re.sub(r"\s+", " ", narration or "").strip()
    if not text:
        return "Unknown"
    if re.match(r"^(upi|imps|neft|rtgs|pos|ach|nach)\b", text, re.I) and "/" in text:
        for tok in text.split("/"):
            tok = tok.strip()
            if (
                len(tok) >= 3
                and tok.lower() not in NOISE_TOKENS
                and not re.fullmatch(r"[\d\W]+", tok)
                and "@" not in tok
                and not re.fullmatch(r"[A-Za-z]{4}\d{6,}\w*", tok)
            ):
                return tok[:80].title() if tok.isupper() else tok[:80]
    return text[:80]


def parse_statement(filename: str, content: bytes, max_rows: int = 5000) -> ParseResult:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        table = _read_csv(content)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        table = _read_xlsx(content)
    elif name.endswith(".xls"):
        raise ParseError("Old .xls files aren't supported. Please re-save the statement as .xlsx or .csv.")
    elif name.endswith(".pdf"):
        raise ParseError("PDF statements aren't supported yet. Please upload a CSV or XLSX export.")
    else:
        raise ParseError("Unsupported file type. Please upload a .csv or .xlsx file.")

    table = [r for r in table if any(str(c or "").strip() for c in r)]
    if not table:
        raise ParseError("The file is empty.")

    # find the header row: must have a date column and some amount-ish column
    header_idx = -1
    for i, row in enumerate(table[:40]):
        h = [_norm(c) for c in row]
        has_date = _find(h, ["date"]) != -1
        has_amt = _find(h, ["amount", "amt", "debit", "credit", "withdrawal", "deposit"]) != -1
        if has_date and has_amt:
            header_idx = i
            break
    if header_idx == -1:
        raise ParseError("Couldn't find a header row with a Date column and Amount/Debit/Credit columns.")

    headers = [_norm(c) for c in table[header_idx]]
    date_i = _find(headers, ["txn date", "transaction date", "tran date", "posting date", "date"], exclude=("value date",))
    if date_i == -1:
        date_i = _find(headers, ["date"])
    merch_i = _find(headers, MERCHANT_KW)
    debit_i = _find(headers, ["withdrawal", "debit", "paid out"], exclude=("balance",))
    credit_i = _find(headers, ["deposit", "credit", "paid in"], exclude=("balance",))
    amount_i = -1
    if debit_i == -1 and credit_i == -1:
        amount_i = _find(headers, ["amount", "amt"], exclude=("balance",))
    type_i = _find(headers, ["dr/cr", "cr/dr", "debit/credit", "txn type", "transaction type", "type"])
    mode_i = _find(headers, ["payment method", "mode", "channel"])

    if debit_i == -1 and credit_i == -1 and amount_i == -1:
        raise ParseError("Couldn't find amount columns in this file.")

    data = table[header_idx + 1:]
    if len(data) > max_rows:
        raise ParseError(f"Too many rows ({len(data)}). Maximum is {max_rows} per import.")

    result = ParseResult()

    def cell(row: list[Any], i: int) -> Any:
        return row[i] if 0 <= i < len(row) else None

    for row in data:
        d = parse_date(cell(row, date_i))
        if d is None:
            result.skipped += 1  # footers, opening-balance lines, etc.
            continue

        hint = None
        if debit_i != -1 or credit_i != -1:
            debit, _ = parse_amount(cell(row, debit_i))
            credit, _ = parse_amount(cell(row, credit_i))
            amount = abs(credit) - abs(debit)
        else:
            amount, hint = parse_amount(cell(row, amount_i))
            tval = _norm(cell(row, type_i)) if type_i != -1 else ""
            if hint is None and tval:
                if tval.startswith("dr") or "debit" in tval or "withdraw" in tval:
                    hint = "dr"
                elif tval.startswith("cr") or "credit" in tval or "deposit" in tval:
                    hint = "cr"
            if hint == "dr":
                amount = -abs(amount)
            elif hint == "cr":
                amount = abs(amount)

        if amount == 0:
            result.skipped += 1
            continue

        raw = str(cell(row, merch_i) or "").strip() if merch_i != -1 else ""
        merchant = clean_merchant(raw)
        category, needs_review = guess_category(f"{raw} {merchant}", amount)
        mode_val = str(cell(row, mode_i) or "").strip() if mode_i != -1 else ""
        result.rows.append(
            ParsedRow(
                date=d,
                merchant=merchant,
                amount=round(amount, 2),
                category=category,
                payment_method=guess_payment_method(raw, mode_val or None),
                notes=raw[:500],
                needs_review=needs_review,
            )
        )

    if not result.rows:
        raise ParseError("No transactions could be found in this file.")
    if result.skipped:
        result.warnings.append(f"{result.skipped} row(s) were skipped (no valid date/amount).")
    return result
